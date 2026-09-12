#!/usr/bin/env python
"""Demonstrate interoperability between odfdo (ODS) and openpyxl (XLSX)."""

import os
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Warning: 'openpyxl' library is not available. Skipping recipe.")
    sys.exit(0)

from odfdo import Document

_DOC_SEQUENCE = 621
DATA = Path(__file__).parent / "data"
SOURCE = "store_table.ods"
OUTPUT_DIR = Path(__file__).parent / "recipes_output" / "openpyxl"


def read_source_document() -> Document:
    """Return the source Document."""
    try:
        source = sys.argv[1]
    except IndexError:
        source = DATA / SOURCE
    return Document(source)


def ods_to_openpyxl(document: Document) -> openpyxl.Workbook:
    """Convert an odfdo Document (ODS) into an openpyxl Workbook (XLSX)."""
    wb = openpyxl.Workbook()
    # Remove the default empty sheet created by openpyxl
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    doc_matrices = document.to_dict(orient="matrix")
    for table_name, rows in doc_matrices.items():
        ws = wb.create_sheet(title=table_name)
        for row in rows:
            ws.append(row)
    return wb


def openpyxl_to_ods(wb: openpyxl.Workbook) -> Document:
    """Convert an openpyxl Workbook (XLSX) into an odfdo Document (ODS)."""
    book_dict: dict[str, list[list[object]]] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        book_dict[sheet_name] = rows
    return Document.from_dict(book_dict)


def openpyxl_interoperability(document: Document) -> None:
    """Demonstrate bidirectional conversion between ODS and XLSX."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1. Convert odfdo Document (ODS) -> openpyxl Workbook (XLSX)
    print("--- 1. Convert ODS Document -> openpyxl Workbook ---")
    wb = ods_to_openpyxl(document)
    print("Workbook Sheet Names:", wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        print(f"Sheet '{name}' ({len(rows)} rows): {rows[0]}")

    output_xlsx_path = OUTPUT_DIR / "store_table.xlsx"
    wb.save(output_xlsx_path)
    print(f"\nSaved Excel workbook to: {output_xlsx_path}")

    # 2. Convert openpyxl Workbook (XLSX) back to an odfdo Document (ODS)
    print("\n--- 2. Convert openpyxl Workbook -> ODS Document ---")
    loaded_wb = openpyxl.load_workbook(output_xlsx_path)
    new_doc = openpyxl_to_ods(loaded_wb)

    print("\nProduct Table Markdown (from converted XLSX):")
    print(new_doc.to_markdown()[0].content)

    output_ods_path = OUTPUT_DIR / "converted_from_excel.ods"
    new_doc.save(output_ods_path)
    print(f"\nSaved converted ODS document to: {output_ods_path}")


def main() -> None:
    document = read_source_document()
    openpyxl_interoperability(document)
    test_unit(document)


def test_unit(document: Document) -> None:
    # only for test suite:
    if "ODFDO_TESTING" not in os.environ:
        return

    # 1. Test ODS -> openpyxl Workbook
    wb = ods_to_openpyxl(document)
    assert wb.sheetnames == ["product", "store"]
    ws_product = wb["product"]
    product_rows = list(ws_product.iter_rows(values_only=True))
    assert product_rows[0] == ("reference", "color", "price")
    assert product_rows[1][0] == "ref01"

    # 2. Test openpyxl Workbook -> ODS Document
    doc_converted = openpyxl_to_ods(wb)
    assert [t.name for t in doc_converted.body.tables] == ["product", "store"]
    prod_dict = doc_converted.to_dict()["product"]
    assert prod_dict["reference"] == ["ref01", "ref02", "ref03"]
    assert prod_dict["color"] == ["white", "blue", "red"]


if __name__ == "__main__":
    main()
